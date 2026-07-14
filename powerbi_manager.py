"""
Task Planner Pro
powerbi_manager.py
Integración con Power BI
"""

from datetime import datetime
from openpyxl import load_workbook


class PowerBIManager:

    def __init__(self, excel_manager, metrics_manager):
        self.excel = excel_manager
        self.metrics = metrics_manager

    def update_metrics_sheet(self):
        snapshot = self.metrics.task_metrics()

        wb = load_workbook(self.excel.file)
        ws = wb["Metricas"]

        ws.append([
            snapshot["fecha"],
            snapshot["pendientes"],
            snapshot["en_curso"],
            snapshot["bloqueadas"],
            snapshot["finalizadas"],
            snapshot["retrasadas"],
            snapshot["total"],
            snapshot["avance_global"]
        ])

        wb.save(self.excel.file)

    def refresh_dimensions(self):
        wb = load_workbook(self.excel.file)

        if "Responsables" in wb.sheetnames:
            ws = wb["Responsables"]
            while ws.max_row > 1:
                ws.delete_rows(2)

            for nombre, datos in sorted(self.metrics.owners_metrics().items()):
                ws.append([
                    nombre,
                    datos["total"],
                    datos["finalizadas"]
                ])

        if "Categorias" in wb.sheetnames:
            ws = wb["Categorias"]
            while ws.max_row > 1:
                ws.delete_rows(2)

            for categoria, total in sorted(self.metrics.categories_metrics().items()):
                ws.append([
                    categoria,
                    total
                ])

        wb.save(self.excel.file)

    def refresh_all(self):
        self.update_metrics_sheet()
        self.refresh_dimensions()

    def export_snapshot(self):
        return {
            "generated": datetime.now().isoformat(),
            **self.metrics.build_snapshot()
        }
